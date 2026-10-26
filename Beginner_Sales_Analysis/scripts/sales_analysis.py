"""
Small Business Sales Performance Analysis
Script: sales_analysis.py
Author: ShadowFox Data Analyst Intern
Objective: Perform exploratory data analysis, derive key business metrics,
           generate high-resolution charts, export analysis results,
           and build an automated multi-tab Excel workbook with a dashboard.
"""

import os
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Configure Matplotlib styling
plt.style.use('seaborn-v0_8-whitegrid' if 'seaborn-v0_8-whitegrid' in plt.style.available else 'default')
plt.rcParams['font.family'] = 'sans-serif'
plt.rcParams['font.size'] = 10
plt.rcParams['axes.titlesize'] = 12
plt.rcParams['axes.titleweight'] = 'bold'
plt.rcParams['axes.labelsize'] = 10
plt.rcParams['axes.labelweight'] = 'bold'
plt.rcParams['figure.titlesize'] = 14
plt.rcParams['figure.titleweight'] = 'bold'

PRIMARY_COLOR = '#1E3E62'
SECONDARY_COLOR = '#0B192C'
ACCENT_COLOR = '#FF6500'
NEUTRAL_COLOR = '#E0E7F1'
SUCCESS_COLOR = '#059669'
ALERT_COLOR = '#DC2626'

def run_sales_analysis(base_dir):
    data_dir = os.path.join(base_dir, "data")
    analysis_dir = os.path.join(base_dir, "analysis")
    charts_dir = os.path.join(base_dir, "charts")
    dashboard_dir = os.path.join(base_dir, "dashboard")
    
    os.makedirs(analysis_dir, exist_ok=True)
    os.makedirs(charts_dir, exist_ok=True)
    os.makedirs(dashboard_dir, exist_ok=True)
    
    cleaned_data_path = os.path.join(data_dir, "cleaned_sales_data.csv")
    raw_data_path = os.path.join(data_dir, "raw_sales_data.csv")
    
    df = pd.read_csv(cleaned_data_path)
    df['Order Date'] = pd.to_datetime(df['Order Date'])
    
    print("=" * 60)
    print("EXECUTING SALES EXPLORATORY DATA ANALYSIS")
    print("=" * 60)
    
    # -------------------------------------------------------------
    # 1. CORE KPI CALCULATIONS
    # -------------------------------------------------------------
    total_revenue = df['Sales'].sum()
    total_cost = df['Cost'].sum()
    total_profit = df['Profit'].sum()
    total_orders = df['Order ID'].nunique()
    total_units_sold = df['Quantity'].sum()
    avg_order_value = total_revenue / total_orders
    avg_profit_per_order = total_profit / total_orders
    overall_profit_margin = (total_profit / total_revenue) * 100
    
    kpi_dict = {
        "Metric": [
            "Total Revenue ($)",
            "Total Cost ($)",
            "Total Profit ($)",
            "Total Orders",
            "Total Units Sold",
            "Average Order Value (AOV) ($)",
            "Average Profit Per Order ($)",
            "Overall Profit Margin (%)"
        ],
        "Value": [
            round(total_revenue, 2),
            round(total_cost, 2),
            round(total_profit, 2),
            total_orders,
            total_units_sold,
            round(avg_order_value, 2),
            round(avg_profit_per_order, 2),
            round(overall_profit_margin, 2)
        ]
    }
    df_kpi = pd.DataFrame(kpi_dict)
    print("\n--- CORE EXECUTIVE KPIS ---")
    print(df_kpi.to_string(index=False))
    
    # -------------------------------------------------------------
    # 2. MONTHLY SALES & PROFIT TRENDS
    # -------------------------------------------------------------
    df_monthly = df.groupby(['Year-Month', 'Month Name', 'Month']).agg(
        Total_Sales=('Sales', 'sum'),
        Total_Profit=('Profit', 'sum'),
        Total_Orders=('Order ID', 'nunique'),
        Units_Sold=('Quantity', 'sum')
    ).reset_index().sort_values(by='Month')
    
    df_monthly['Profit_Margin_%'] = ((df_monthly['Total_Profit'] / df_monthly['Total_Sales']) * 100).round(2)
    df_monthly['AOV'] = (df_monthly['Total_Sales'] / df_monthly['Total_Orders']).round(2)
    df_monthly['Sales_MoM_Growth_%'] = df_monthly['Total_Sales'].pct_change() * 100
    df_monthly['Sales_MoM_Growth_%'] = df_monthly['Sales_MoM_Growth_%'].fillna(0).round(2)
    
    # -------------------------------------------------------------
    # 3. CATEGORY PERFORMANCE
    # -------------------------------------------------------------
    df_category = df.groupby('Category').agg(
        Total_Sales=('Sales', 'sum'),
        Total_Profit=('Profit', 'sum'),
        Units_Sold=('Quantity', 'sum'),
        Order_Count=('Order ID', 'count')
    ).reset_index()
    df_category['Profit_Margin_%'] = ((df_category['Total_Profit'] / df_category['Total_Sales']) * 100).round(2)
    df_category['Sales_Share_%'] = ((df_category['Total_Sales'] / total_revenue) * 100).round(2)
    df_category = df_category.sort_values(by='Total_Sales', ascending=False)
    
    # -------------------------------------------------------------
    # 4. REGIONAL PERFORMANCE
    # -------------------------------------------------------------
    df_region = df.groupby('Region').agg(
        Total_Sales=('Sales', 'sum'),
        Total_Profit=('Profit', 'sum'),
        Total_Orders=('Order ID', 'nunique'),
        Units_Sold=('Quantity', 'sum')
    ).reset_index()
    df_region['Profit_Margin_%'] = ((df_region['Total_Profit'] / df_region['Total_Sales']) * 100).round(2)
    df_region['AOV'] = (df_region['Total_Sales'] / df_region['Total_Orders']).round(2)
    df_region = df_region.sort_values(by='Total_Sales', ascending=False)
    
    # -------------------------------------------------------------
    # 5. PRODUCT PERFORMANCE (TOP & BOTTOM)
    # -------------------------------------------------------------
    df_product = df.groupby(['Product', 'Category']).agg(
        Total_Sales=('Sales', 'sum'),
        Total_Profit=('Profit', 'sum'),
        Units_Sold=('Quantity', 'sum'),
        Order_Count=('Order ID', 'count')
    ).reset_index()
    df_product['Profit_Margin_%'] = ((df_product['Total_Profit'] / df_product['Total_Sales']) * 100).round(2)
    df_product = df_product.sort_values(by='Total_Sales', ascending=False)
    
    top_5_products = df_product.head(5).copy()
    bottom_5_products = df_product.tail(5).copy()
    
    # -------------------------------------------------------------
    # 6. PAYMENT METHOD ANALYSIS
    # -------------------------------------------------------------
    df_payment = df.groupby('Payment Method').agg(
        Total_Sales=('Sales', 'sum'),
        Total_Profit=('Profit', 'sum'),
        Transactions=('Order ID', 'count')
    ).reset_index()
    df_payment['Sales_Share_%'] = ((df_payment['Total_Sales'] / total_revenue) * 100).round(2)
    df_payment = df_payment.sort_values(by='Total_Sales', ascending=False)
    
    # -------------------------------------------------------------
    # 7. EXPORT ANALYSIS RESULTS TO CSV
    # -------------------------------------------------------------
    df_kpi.to_csv(os.path.join(analysis_dir, "analysis_results.csv"), index=False)
    df_monthly.to_csv(os.path.join(analysis_dir, "monthly_sales_summary.csv"), index=False)
    df_category.to_csv(os.path.join(analysis_dir, "category_sales_summary.csv"), index=False)
    df_region.to_csv(os.path.join(analysis_dir, "regional_sales_summary.csv"), index=False)
    df_product.to_csv(os.path.join(analysis_dir, "product_performance_summary.csv"), index=False)
    df_payment.to_csv(os.path.join(analysis_dir, "payment_method_summary.csv"), index=False)
    print("\nAll analytical summaries successfully exported to 'analysis/' directory.")
    
    # -------------------------------------------------------------
    # 8. VISUALIZATIONS GENERATION
    # -------------------------------------------------------------
    print("\nGenerating high-resolution visualization charts...")
    
    # Chart 1: Monthly Sales & Profit Trend
    fig, ax1 = plt.subplots(figsize=(11, 5.5), dpi=300)
    months = df_monthly['Month Name']
    x = np.arange(len(months))
    
    bars = ax1.bar(x - 0.2, df_monthly['Total_Sales'], width=0.4, label='Revenue ($)', color=PRIMARY_COLOR, alpha=0.9)
    bars2 = ax1.bar(x + 0.2, df_monthly['Total_Profit'], width=0.4, label='Profit ($)', color=ACCENT_COLOR, alpha=0.9)
    
    ax1.set_xlabel('Month', fontweight='bold', labelpad=8)
    ax1.set_ylabel('Amount ($)', fontweight='bold', labelpad=8)
    ax1.set_xticks(x)
    ax1.set_xticklabels(months)
    ax1.yaxis.set_major_formatter('${x:,.0f}')
    ax1.set_title('Small Business Monthly Revenue & Profit Trends (2024)', fontsize=14, pad=15)
    ax1.legend(loc='upper left', frameon=True)
    
    # Add value labels on top of revenue bars
    for bar in bars:
        height = bar.get_height()
        ax1.annotate(f'${height/1000:.1f}k',
                     xy=(bar.get_x() + bar.get_width() / 2, height),
                     xytext=(0, 3), textcoords="offset points",
                     ha='center', va='bottom', fontsize=8, fontweight='bold')
                     
    plt.tight_layout()
    chart1_path = os.path.join(charts_dir, "monthly_sales_trend.png")
    fig.savefig(chart1_path, dpi=300)
    plt.close()
    
    # Chart 2: Category Performance (Revenue vs Profit Margin)
    fig, ax1 = plt.subplots(figsize=(10, 5), dpi=300)
    ax2 = ax1.twinx()
    
    cats = df_category['Category']
    x_cat = np.arange(len(cats))
    
    rects = ax1.bar(x_cat, df_category['Total_Sales'], width=0.45, color='#2563EB', alpha=0.85, label='Total Sales ($)')
    line = ax2.plot(x_cat, df_category['Profit_Margin_%'], color='#EA580C', marker='o', linewidth=2.5, markersize=8, label='Profit Margin (%)')
    
    ax1.set_xticks(x_cat)
    ax1.set_xticklabels(cats, fontweight='bold')
    ax1.set_ylabel('Total Sales ($)', fontweight='bold', color='#1E3E62')
    ax2.set_ylabel('Profit Margin (%)', fontweight='bold', color='#EA580C')
    ax1.yaxis.set_major_formatter('${x:,.0f}')
    ax2.yaxis.set_major_formatter('{x:.1f}%')
    ax2.grid(False)
    
    for r in rects:
        h = r.get_height()
        ax1.annotate(f'${h:,.0f}', xy=(r.get_x() + r.get_width()/2, h), xytext=(0, 3),
                     textcoords="offset points", ha='center', va='bottom', fontsize=9, fontweight='bold')
                     
    for i, val in enumerate(df_category['Profit_Margin_%']):
        ax2.annotate(f'{val:.1f}%', xy=(x_cat[i], val), xytext=(0, 8),
                     textcoords="offset points", ha='center', va='bottom', fontsize=9, fontweight='bold', color='#EA580C')
                     
    plt.title('Category Performance: Revenue vs Profit Margin', fontsize=13, pad=15)
    plt.tight_layout()
    chart2_path = os.path.join(charts_dir, "category_sales_profit.png")
    fig.savefig(chart2_path, dpi=300)
    plt.close()
    
    # Chart 3: Regional Sales & Profit Comparison
    fig, ax = plt.subplots(figsize=(10, 5), dpi=300)
    df_reg_sorted = df_region.sort_values(by='Total_Sales', ascending=True)
    y_pos = np.arange(len(df_reg_sorted))
    
    ax.barh(y_pos - 0.18, df_reg_sorted['Total_Sales'], height=0.35, label='Revenue ($)', color='#0284C7', alpha=0.9)
    ax.barh(y_pos + 0.18, df_reg_sorted['Total_Profit'], height=0.35, label='Profit ($)', color='#10B981', alpha=0.9)
    
    ax.set_yticks(y_pos)
    ax.set_yticklabels(df_reg_sorted['Region'], fontweight='bold')
    ax.xaxis.set_major_formatter('${x:,.0f}')
    ax.set_xlabel('Amount ($)', fontweight='bold')
    ax.set_title('Regional Sales and Profitability Comparison', fontsize=13, pad=15)
    ax.legend(loc='lower right', frameon=True)
    
    for i, row in df_reg_sorted.reset_index().iterrows():
        ax.text(row['Total_Sales'] + 500, i - 0.18, f"${row['Total_Sales']:,.0f}", va='center', fontsize=8, fontweight='bold', color='#0369A1')
        ax.text(row['Total_Profit'] + 500, i + 0.18, f"${row['Total_Profit']:,.0f}", va='center', fontsize=8, fontweight='bold', color='#047857')
        
    plt.tight_layout()
    chart3_path = os.path.join(charts_dir, "regional_performance.png")
    fig.savefig(chart3_path, dpi=300)
    plt.close()
    
    # Chart 4: Top 5 and Bottom 5 Products by Revenue
    fig, (ax_top, ax_bot) = plt.subplots(1, 2, figsize=(14, 5.5), dpi=300)
    
    # Top 5
    top_sorted = top_5_products.sort_values(by='Total_Sales', ascending=True)
    ax_top.barh(top_sorted['Product'], top_sorted['Total_Sales'], color='#15803D', alpha=0.85)
    ax_top.set_title('Top 5 Products by Revenue', fontsize=12, fontweight='bold')
    ax_top.xaxis.set_major_formatter('${x:,.0f}')
    for i, v in enumerate(top_sorted['Total_Sales']):
        ax_top.text(v + 300, i, f'${v:,.0f}', va='center', fontsize=8, fontweight='bold')
        
    # Bottom 5
    bot_sorted = bottom_5_products.sort_values(by='Total_Sales', ascending=True)
    ax_bot.barh(bot_sorted['Product'], bot_sorted['Total_Sales'], color='#E11D48', alpha=0.85)
    ax_bot.set_title('Lowest 5 Products by Revenue', fontsize=12, fontweight='bold')
    ax_bot.xaxis.set_major_formatter('${x:,.0f}')
    for i, v in enumerate(bot_sorted['Total_Sales']):
        ax_bot.text(v + 200, i, f'${v:,.0f}', va='center', fontsize=8, fontweight='bold')
        
    plt.tight_layout()
    chart4_path = os.path.join(charts_dir, "top_bottom_products.png")
    fig.savefig(chart4_path, dpi=300)
    plt.close()
    
    # Chart 5: Payment Method Distribution
    fig, ax = plt.subplots(figsize=(8, 5), dpi=300)
    colors = ['#3B82F6', '#10B981', '#F59E0B', '#8B5CF6', '#EC4899', '#6B7280']
    wedges, texts, autotexts = ax.pie(
        df_payment['Total_Sales'],
        labels=df_payment['Payment Method'],
        autopct='%1.1f%%',
        startangle=140,
        colors=colors[:len(df_payment)],
        wedgeprops=dict(width=0.4, edgecolor='w', linewidth=2)
    )
    plt.setp(autotexts, size=9, weight="bold")
    plt.setp(texts, size=10, weight="bold")
    ax.set_title('Revenue Breakdown by Payment Method', fontsize=13, pad=15)
    plt.tight_layout()
    chart5_path = os.path.join(charts_dir, "payment_method_distribution.png")
    fig.savefig(chart5_path, dpi=300)
    plt.close()
    
    print("Visual charts saved to:", charts_dir)
    
    # -------------------------------------------------------------
    # 9. BUILD MULTI-TAB EXCEL WORKBOOK WITH DASHBOARD
    # -------------------------------------------------------------
    print("\nBuilding professional Excel Workbook: Small_Business_Sales_Dashboard.xlsx...")
    excel_path = os.path.join(dashboard_dir, "Small_Business_Sales_Dashboard.xlsx")
    
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    # Common styles
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    title_font = Font(name='Segoe UI', size=16, bold=True, color='0B192C')
    sub_title_font = Font(name='Segoe UI', size=11, italic=True, color='555555')
    bold_font = Font(name='Segoe UI', size=10, bold=True)
    normal_font = Font(name='Segoe UI', size=10)
    kpi_num_font = Font(name='Segoe UI', size=18, bold=True, color='1E3E62')
    kpi_label_font = Font(name='Segoe UI', size=9, bold=True, color='4B5563')
    
    header_fill = PatternFill(start_color='1E3E62', end_color='1E3E62', fill_type='solid')
    sub_header_fill = PatternFill(start_color='334155', end_color='334155', fill_type='solid')
    kpi_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    accent_kpi_fill = PatternFill(start_color='EFF6FF', end_color='EFF6FF', fill_type='solid')
    
    thin_border_side = Side(style='thin', color='D1D5DB')
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(bottom=Side(style='medium', color='1E3E62'))
    
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')
    
    # ------------------ SHEET 1: RAW DATA ------------------
    ws_raw = wb.create_sheet(title="Raw Data")
    df_raw = pd.read_csv(raw_data_path)
    
    # Header
    ws_raw.append(list(df_raw.columns))
    for col_idx in range(1, len(df_raw.columns) + 1):
        cell = ws_raw.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = sub_header_fill
        cell.alignment = center_align
        
    for row in df_raw.itertuples(index=False):
        ws_raw.append(list(row))
        
    for col in ws_raw.columns:
        col_letter = get_column_letter(col[0].column)
        ws_raw.column_dimensions[col_letter].width = 16
        
    # ------------------ SHEET 2: CLEANED DATA ------------------
    ws_clean = wb.create_sheet(title="Cleaned Data")
    clean_cols = list(df.columns)
    ws_clean.append(clean_cols)
    
    for col_idx in range(1, len(clean_cols) + 1):
        cell = ws_clean.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        
    for row in df.itertuples(index=False):
        ws_clean.append(list(row))
        
    # Format Currency and Numbers in Cleaned Data sheet
    num_rows_clean = len(df) + 1
    for r in range(2, num_rows_clean + 1):
        ws_clean.cell(row=r, column=1).alignment = center_align # Order ID
        ws_clean.cell(row=r, column=2).alignment = center_align # Order Date
        ws_clean.cell(row=r, column=5).number_format = '#,##0'  # Quantity
        ws_clean.cell(row=r, column=6).number_format = '$#,##0.00' # Unit Price
        ws_clean.cell(row=r, column=7).number_format = '$#,##0.00' # Sales
        ws_clean.cell(row=r, column=8).number_format = '$#,##0.00' # Cost
        ws_clean.cell(row=r, column=9).number_format = '$#,##0.00' # Profit
        ws_clean.cell(row=r, column=12).number_format = '0.00"%"' # Profit Margin
        
    for col in ws_clean.columns:
        col_letter = get_column_letter(col[0].column)
        ws_clean.column_dimensions[col_letter].width = 16
        
    # ------------------ SHEET 3: ANALYSIS ------------------
    ws_analysis = wb.create_sheet(title="Analysis")
    ws_analysis.views.sheetView[0].showGridLines = True
    
    ws_analysis["A1"] = "SMALL BUSINESS SALES ANALYSIS - SUMMARY CALCULATIONS"
    ws_analysis["A1"].font = title_font
    ws_analysis["A2"] = "Dynamic Excel Formulas operating directly on the Cleaned Data table"
    ws_analysis["A2"].font = sub_title_font
    
    # KPI Formulas block
    ws_analysis["A4"] = "EXECUTIVE METRIC"
    ws_analysis["B4"] = "FORMULA EXPRESSION"
    ws_analysis["C4"] = "CALCULATED VALUE"
    
    for c in ["A4", "B4", "C4"]:
        ws_analysis[c].font = header_font
        ws_analysis[c].fill = header_fill
        ws_analysis[c].alignment = center_align
        
    metrics_formulas = [
        ("Total Sales Revenue", "=SUM('Cleaned Data'!G2:G" + str(num_rows_clean) + ")", "$#,##0.00"),
        ("Total Production Cost", "=SUM('Cleaned Data'!H2:H" + str(num_rows_clean) + ")", "$#,##0.00"),
        ("Total Net Profit", "=SUM('Cleaned Data'!I2:I" + str(num_rows_clean) + ")", "$#,##0.00"),
        ("Total Transactions / Orders", "=COUNTA('Cleaned Data'!A2:A" + str(num_rows_clean) + ")", "#,##0"),
        ("Total Units Sold", "=SUM('Cleaned Data'!E2:E" + str(num_rows_clean) + ")", "#,##0"),
        ("Average Order Value (AOV)", "=C5/C8", "$#,##0.00"),
        ("Overall Profit Margin (%)", "=(C7/C5)*100", "0.00\"%\"")
    ]
    
    for idx, (lbl, form, fmt) in enumerate(metrics_formulas, start=5):
        ws_analysis[f"A{idx}"] = lbl
        ws_analysis[f"A{idx}"].font = bold_font
        ws_analysis[f"B{idx}"] = form.replace("=", "Formula: ")
        ws_analysis[f"B{idx}"].font = normal_font
        ws_analysis[f"C{idx}"] = form
        ws_analysis[f"C{idx}"].font = bold_font
        ws_analysis[f"C{idx}"].number_format = fmt
        
        for col_let in ["A", "B", "C"]:
            ws_analysis[f"{col_let}{idx}"].border = thin_border
            
    # Category summary table in Analysis Sheet using SUMIFS formulas
    ws_analysis["A14"] = "CATEGORY PERFORMANCE BREAKDOWN (FORMULA DRIVEN)"
    ws_analysis["A14"].font = Font(name='Segoe UI', size=12, bold=True, color='1E3E62')
    
    cat_headers = ["Category", "Formula Revenue ($)", "Formula Profit ($)", "Profit Margin %"]
    for c_idx, h in enumerate(cat_headers, start=1):
        cell = ws_analysis.cell(row=15, column=c_idx)
        cell.value = h
        cell.font = header_font
        cell.fill = sub_header_fill
        cell.alignment = center_align
        
    categories_list = sorted(df['Category'].unique())
    for r_idx, cat in enumerate(categories_list, start=16):
        ws_analysis.cell(row=r_idx, column=1, value=cat).font = bold_font
        ws_analysis.cell(row=r_idx, column=2, value=f"=SUMIFS('Cleaned Data'!$G$2:$G${num_rows_clean}, 'Cleaned Data'!$D$2:$D${num_rows_clean}, A{r_idx})").number_format = '$#,##0.00'
        ws_analysis.cell(row=r_idx, column=3, value=f"=SUMIFS('Cleaned Data'!$I$2:$I${num_rows_clean}, 'Cleaned Data'!$D$2:$D${num_rows_clean}, A{r_idx})").number_format = '$#,##0.00'
        ws_analysis.cell(row=r_idx, column=4, value=f"=IFERROR((C{r_idx}/B{r_idx})*100, 0)").number_format = '0.00"%"'
        
        for c in range(1, 5):
            ws_analysis.cell(row=r_idx, column=c).border = thin_border
            
    # Auto-adjust column widths for Analysis sheet
    ws_analysis.column_dimensions['A'].width = 30
    ws_analysis.column_dimensions['B'].width = 32
    ws_analysis.column_dimensions['C'].width = 24
    ws_analysis.column_dimensions['D'].width = 20

    # ------------------ SHEET 4: DASHBOARD ------------------
    ws_dash = wb.create_sheet(title="Dashboard")
    ws_dash.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws_dash.merge_cells("A1:K1")
    ws_dash["A1"] = "SMALL BUSINESS EXECUTIVE SALES DASHBOARD (FY 2024)"
    ws_dash["A1"].font = Font(name='Segoe UI', size=18, bold=True, color='FFFFFF')
    ws_dash["A1"].fill = PatternFill(start_color='0B192C', end_color='0B192C', fill_type='solid')
    ws_dash["A1"].alignment = Alignment(horizontal='center', vertical='center')
    ws_dash.row_dimensions[1].height = 40
    
    ws_dash.merge_cells("A2:K2")
    ws_dash["A2"] = "Performance Monitoring | Monthly Revenue, Category Margins, Regional Profitability & Product Sales"
    ws_dash["A2"].font = Font(name='Segoe UI', size=10, italic=True, color='475569')
    ws_dash["A2"].alignment = Alignment(horizontal='center', vertical='center')
    ws_dash.row_dimensions[2].height = 20
    
    # KPI Cards (Row 4 to 6)
    kpis = [
        ("TOTAL REVENUE", f"=Analysis!C5", "$#,##0", "A", "B"),
        ("TOTAL PROFIT", f"=Analysis!C7", "$#,##0", "C", "D"),
        ("OVERALL MARGIN", f"=Analysis!C11", "0.0\"%\"", "E", "F"),
        ("TOTAL ORDERS", f"=Analysis!C8", "#,##0", "G", "H"),
        ("AVG ORDER VALUE", f"=Analysis!C10", "$#,##0.00", "I", "J")
    ]
    
    ws_dash.row_dimensions[4].height = 18
    ws_dash.row_dimensions[5].height = 30
    
    for title, formula, fmt, c1, c2 in kpis:
        # Header cell
        ws_dash.merge_cells(f"{c1}4:{c2}4")
        card_lbl = ws_dash[f"{c1}4"]
        card_lbl.value = title
        card_lbl.font = kpi_label_font
        card_lbl.fill = kpi_fill
        card_lbl.alignment = center_align
        
        # Value cell
        ws_dash.merge_cells(f"{c1}5:{c2}5")
        card_val = ws_dash[f"{c1}5"]
        card_val.value = formula
        card_val.font = kpi_num_font
        card_val.fill = accent_kpi_fill
        card_val.number_format = fmt
        card_val.alignment = center_align
        
        # Borders
        for r in range(4, 6):
            for col_letter in [c1, c2]:
                ws_dash[f"{col_letter}{r}"].border = thin_border
                
    # Monthly Trends Summary Table
    ws_dash["A8"] = "MONTHLY SALES PERFORMANCE SUMMARY"
    ws_dash["A8"].font = Font(name='Segoe UI', size=11, bold=True, color='1E3E62')
    
    m_headers = ["Month", "Revenue ($)", "Profit ($)", "Orders", "Margin %"]
    for c_i, h in enumerate(m_headers, start=1):
        cell = ws_dash.cell(row=9, column=c_i, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        
    for r_i, m_row in df_monthly.iterrows():
        row_num = 10 + r_i
        ws_dash.cell(row=row_num, column=1, value=m_row['Month Name']).font = bold_font
        ws_dash.cell(row=row_num, column=2, value=m_row['Total_Sales']).number_format = '$#,##0'
        ws_dash.cell(row=row_num, column=3, value=m_row['Total_Profit']).number_format = '$#,##0'
        ws_dash.cell(row=row_num, column=4, value=m_row['Total_Orders']).number_format = '#,##0'
        ws_dash.cell(row=row_num, column=5, value=m_row['Profit_Margin_%'] / 100.0).number_format = '0.0%'
        
        for c in range(1, 6):
            ws_dash.cell(row=row_num, column=c).border = thin_border
            
    # Add embedded chart image to the dashboard!
    img = openpyxl.drawing.image.Image(chart1_path)
    img.width = 580
    img.height = 290
    ws_dash.add_image(img, "G8")
    
    # Add category chart image below
    img2 = openpyxl.drawing.image.Image(chart2_path)
    img2.width = 540
    img2.height = 270
    ws_dash.add_image(img2, "A24")
    
    # Add regional chart image beside
    img3 = openpyxl.drawing.image.Image(chart3_path)
    img3.width = 540
    img3.height = 270
    ws_dash.add_image(img3, "G24")
    
    # Adjust column dimensions
    ws_dash.column_dimensions['A'].width = 14
    ws_dash.column_dimensions['B'].width = 16
    ws_dash.column_dimensions['C'].width = 16
    ws_dash.column_dimensions['D'].width = 14
    ws_dash.column_dimensions['E'].width = 14
    ws_dash.column_dimensions['F'].width = 14
    ws_dash.column_dimensions['G'].width = 16
    ws_dash.column_dimensions['H'].width = 16
    ws_dash.column_dimensions['I'].width = 16
    ws_dash.column_dimensions['J'].width = 16
    ws_dash.column_dimensions['K'].width = 16
    
    wb.save(excel_path)
    print(f"Excel workbook successfully generated and saved to: {excel_path}")
    print("=" * 60)
    print("SALES ANALYSIS AND DASHBOARD GENERATION COMPLETE")
    print("=" * 60)
    
    return {
        "kpis": df_kpi,
        "monthly": df_monthly,
        "category": df_category,
        "region": df_region,
        "top_products": top_5_products,
        "bottom_products": bottom_5_products,
        "payment": df_payment
    }

if __name__ == "__main__":
    base_directory = r"C:\Users\keert\.gemini\antigravity-ide\scratch\ShadowFox_Data_Analyst_Internship\Beginner_Sales_Analysis"
    run_sales_analysis(base_directory)
